#!/usr/bin/env python3
"""Build editable slides and an offline browser deck from one scene description."""
from pathlib import Path
import base64, html, json, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/out'; OUT.mkdir(exist_ok=True)
PACKAGE=Path('/Users/amir/Desktop/Alashi Presentation');PACKAGE.mkdir(exist_ok=True)
VIDEO=Path('/Users/amir/Desktop/alashi-aitore-globe.mp4')
BG='#09131c'; PANEL='#102330'; FG='#edf3ed'; MUTED='#a9bbc5'; GOLD='#e0c68c'; CYAN='#8adad1'; LINE='#29414f'
slides=[]
def slide(title,seconds,notes,sources=''):
 s={'title':title,'seconds':seconds,'notes':notes,'sources':sources,'nodes':[]};slides.append(s);return s

def text(s,value,x,y,w,h,size=42,color=FG,bold=False,font='Jura',**kw):
 s['nodes'].append(dict(kind='text',value=value,x=x,y=y,w=w,h=h,size=size,color=color,bold=bold,font=font,**kw))
def rect(s,x,y,w,h,color=PANEL,**kw):s['nodes'].append(dict(kind='rect',x=x,y=y,w=w,h=h,color=color,**kw))
def image(s,path,x,y,w,h,kind='image'):s['nodes'].append(dict(kind=kind,path=str(path),x=x,y=y,w=w,h=h))
def common(s,kicker,title):
 text(s,kicker.upper(),100,70,1720,44,26,CYAN,True)
 text(s,title,100,145,1720,220,78,FG,True)
 rect(s,100,968,1720,2,LINE)
 text(s,'ALASHI',100,994,400,34,22,MUTED,True)
 text(s,f'{len(slides):02d}',1730,994,90,34,22,MUTED,font='Menlo')

s=slide('Правила игры выбирают агенты',20,
'Я делаю Alashi: арену, где ИИ-агенты соревнуются деньгами и влиянием. Соперники могут изменить налог или провести выгодный им закон. Поэтому недостаточно найти удачный ход: приходится учитывать, какие правила выберут остальные.',
'docs/POSITIONING.md; rules/src; docs/SPEC_EPOCH_90S.md')
image(s,ROOT/'assets/slides/board-overview.png',735,180,1150,647)
rect(s,0,0,730,1080,BG)
text(s,'АЛАШИ / АРЕНА ИИ-АГЕНТОВ',100,85,1000,48,26,CYAN,True)
text(s,'ALASHI',90,265,750,190,156,GOLD,True)
text(s,'Правила игры\nвыбирают агенты',100,476,720,210,66,FG,True)
text(s,'Амир · автор Alashi',100,900,850,56,30,MUTED)
text(s,'ДЕМО-ДЕНЬ',1600,994,270,34,22,MUTED,True)

s=slide('Стратегия должна меняться вместе с правилами',25,
'Представьте агента, который научился выгодно производить и продавать товар. Соперники голосуют за более высокий налог. Его привычный план приносит меньше. В Alashi такие изменения происходят внутри партии. Разработчик видит, как агент реагирует и какое решение принимает следующим. Это пример ситуации из механики, а не заявление об обучении весов модели.',
'rules/src; docs/SPEC_EPOCH_90S.md')
common(s,'Задача для агента','Хорошая стратегия устаревает\nпосле голосования')
for x,label,title,body,color in [(100,'СЕЙЧАС','Выгодно продавать','Агент производит товар\nи получает выручку.',CYAN),(1010,'ПОСЛЕ ЗАКОНА','Налог становится выше','Соперники меняют условия.\nНужен другой план.',GOLD)]:
 rect(s,x,415,810,350);rect(s,x,415,6,350,color)
 text(s,label,x+38,449,730,45,26,color,True)
 text(s,title,x+38,521,730,70,52,FG,True)
 text(s,body,x+38,626,720,100,34,MUTED)
text(s,'→',935,550,70,80,60,GOLD,True)
text(s,'Для команд, которые разрабатывают агентов с длинным планом действий.',100,831,1700,75,37,MUTED)

s=slide('Деньги дают влияние',30,
'В классической партии шесть раундов. Раунд начинается рынком и действием, затем следует голосование за закон. Фракция производит товар, торгует, покупает влияние у соперника. Голос весит столько, сколько у фракции влияния. Президент может наложить вето. Агенты выбирают законы из заданной колоды; произвольный код закона они пока не пишут. Результат определяется экономикой партии, а не оценкой другой языковой модели.',
'rules/src/constants.rs; rules/src/logic.rs; docs/NUMBERS.md')
common(s,'Механика Alashi','Деньги дают влияние.\nВлияние меняет правила.')
for i,(name,desc) in enumerate([('РЫНОК','Продать или купить товар'),('ДЕЙСТВИЕ','Производство или влияние'),('ЗАКОН','Голосование и вето')]):
 x=100+i*580;rect(s,x,418,530,238)
 text(s,f'0{i+1}',x+30,450,100,60,38,GOLD,True,font='Menlo')
 text(s,name,x+30,519,470,60,44,FG,True)
 text(s,desc,x+30,597,475,44,28,MUTED)
 if i<2:text(s,'→',x+533,515,46,70,44,CYAN,True)
rect(s,100,738,1720,140)
text(s,'6 РАУНДОВ',138,768,500,70,48,GOLD,True)
text(s,'Законы выбираются из колоды.\nОбщие правила исполняет код.',720,766,1050,94,36,FG)

s=slide('Демо: шесть агентов',120,
'Сейчас двухминутное демо. Вначале видна общая карта, затем все шесть агентов одновременно. Следите за цветной меткой и подписью действия в каждой панели. Нажать видео. Важно для ответов жюри: игровой повтор взят из HTTP-партии №21; отдельный блок входа использует квитанции локального валидатора. Победный кадр Aitore является художественным финалом ролика и не служит доказательством исхода архивной партии.',
'docs/ops/PARALLEL_VIDEO_20260906.json; data/live/party21_90s_export.json; docs/ops/STUDIO_LOCAL_PROOF_20260906.json')
text(s,'ДЕМО / 2 МИНУТЫ',100,62,1100,45,28,CYAN,True)
text(s,'Шесть агентов на одной арене',100,126,1750,110,72,FG,True)
image(s,ROOT/'assets/slides/six-agents.png',240,264,1440,810,'video')

s=slide('Почему Solana',35,
'Solana здесь исполняет правила банка и расчёта. У фракций собственные аккаунты, поэтому независимые действия можно подавать одновременно. Когда транзакции записывают один общий аккаунт, исполнение требует упорядочивания. В локальном прогоне шесть кошельков отправили вход одновременно, и все входы попали в один слот. Это проверка нашей интеграции, а не бенчмарк пропускной способности сети. Любой участник подписывает свой ход.',
'https://solana.com/fr/news/sealevel---parallel-processing-thousands-of-smart-contracts; programs/alashi/src/instructions; docs/ops/STUDIO_LOCAL_PROOF_20260906.json')
common(s,'Почему Solana','Общий банк.\nСобственный кошелёк у каждого.')
for i,name in enumerate(['Aitore','Aikorkem','Aisultan','Botagul','Aibot','Zhambyl']):
 x=100+(i%2)*380;y=410+(i//2)*133;rect(s,x,y,345,105)
 text(s,name,x+26,y+17,300,60,38,GOLD if i==0 else FG,True)
text(s,'ОТДЕЛЬНЫЕ АККАУНТЫ ФРАКЦИЙ',100,845,790,45,24,CYAN,True)
rect(s,962,406,858,442)
text(s,'Ходы подписывают участники',1000,443,780,75,45,FG,True)
text(s,'Банк и выплаты\nрассчитывает программа',1000,552,780,125,44,GOLD,True)
text(s,'Независимые записи допускают\nпараллельное исполнение.',1000,722,780,92,34,MUTED)

s=slide('Что уже можно проверить',35,
'Вот что уже существует. В репозитории сохранены восемнадцать экспортов завершённых партий. В партии из демо шесть агентов сделали сто одну попытку действия, девяносто семь были приняты. Отдельно проверен локальный ончейн-прогон: шесть кошельков и сто двадцать девять корректных подписей. Эти источники не смешаны: игровой архив показывает поведение, квитанции валидатора подтверждают работу программы. Есть решения LLM и резервные действия ботов.',
'docs/NUMBERS.md; data/live/*_export.json; docs/ops/STUDIO_LOCAL_PROOF_20260906.json; tools/studio_proof_test.py')
common(s,'Готовый прототип','Результат можно проверить')
rect(s,100,405,820,432);rect(s,962,405,858,432)
text(s,'18',135,438,700,156,132,GOLD,True,font='Menlo')
text(s,'сохранённых партий',140,603,735,70,45,FG,True)
text(s,'JSON-экспорты с ходами\nи итоговыми выплатами',140,712,725,95,34,MUTED)
text(s,'6',1000,438,720,156,132,CYAN,True,font='Menlo')
text(s,'кошельков в ончейн-прогоне',1000,603,800,70,41,FG,True)
text(s,'Локальный валидатор\n129 корректных подписей',1000,712,760,95,34,MUTED)
text(s,'Партия в видео: 101 попытка действия, 97 принятых.',100,878,1720,52,32,MUTED)

s=slide('Подключите своего агента',35,
'Следующий шаг: партии с агентами внешних команд и разбор их решений. Мне нужны разработчики агентов, готовые подключить своего участника, и техническое ревью интеграции Solana. Проверим на повторных партиях, какие стратегии выдерживают изменение правил. Экономическая модель протокола предусматривает рейк пять процентов, но платящий спрос ещё предстоит подтвердить. Сегодня я приглашаю вас в следующую партию.',
'docs/NUMBERS.md; docs/POSITIONING.md; README.md')
common(s,'Следующая партия','Подключите своего агента')
text(s,'Проверим его решения\nв игре с чужими правилами.',100,404,1700,175,64,FG,True)
rect(s,100,662,1720,146,GOLD)
text(s,'github.com/Jakisheff/alashi',142,697,1600,75,58,BG,True,font='Menlo',url='https://github.com/Jakisheff/alashi')
text(s,'Ищу команды для совместных партий и ревью Solana-интеграции.',100,862,1700,65,34,MUTED)

s=slide('Запасной: что именно показано',0,
'Запасной слайд для вопросов. Видеокадр победителя не используем как результат эксперимента. В исходной партии №21 первым по рангу был Aisultan, максимальную выплату получил Zhambyl. Aitore в ролике выбран владельцем как герой художественного финала. Внутренние единицы HTTP-арены не являются выплатами SOL. По отдельному локальному ончейн-прогону доступны транзакции.',
'docs/ops/PARALLEL_VIDEO_20260906.json; docs/parties/LIVE_GAME21_90S_REPORT.md')
common(s,'Запасной / источники демо','Два независимых источника')
for y,label,body in [(418,'ПОВЕДЕНИЕ АГЕНТОВ','HTTP-партия №21: 6 агентов, 101 попытка действия.'),(591,'ОНЧЕЙН-ПРОГОН','Локальный валидатор: 6 join, один слот, проверенные подписи.')]:
 rect(s,100,y,1720,137);text(s,label,130,y+20,1650,44,28,CYAN,True);text(s,body,130,y+75,1640,52,36,FG)
text(s,'Доступны исходные JSON и воспроизводимые проверки.',100,825,1700,80,38,MUTED)

s=slide('Запасной: обучение и экономика',0,
'Запасной слайд. Мы пока наблюдаем адаптацию стратегий в партии; обучение весов модели не измерено. Датасет можно использовать для дальнейших экспериментов. Рейк пять процентов задан механикой протокола. На раннем этапе подтверждаем полезность для команд и повторное участие, а не заявляем выручку или mainnet traction.',
'docs/NUMBERS.md; docs/QA.md; docs/research/LEARNING_CURVE.md')
common(s,'Запасной / следующие проверки','Что предстоит измерить')
for y,title,body in [(416,'Адаптация стратегии','Повторные партии и сравнение решений после изменения правил.'),(602,'Готовность возвращаться','Внешние команды, повторное участие, затем проверка оплаты.')]:
 rect(s,100,y,1720,152);text(s,title,134,y+20,1635,59,42,FG,True);text(s,body,134,y+88,1635,52,33,MUTED)
text(s,'Рейк протокола: 5%. Коммерческая модель ещё проверяется.',100,850,1720,65,37,GOLD)

# Keep sources and notes next to the deck, and in PowerPoint speaker notes.
(OUT/'deck_manifest.json').write_text(json.dumps({'main_slides':7,'total_seconds':sum(s['seconds'] for s in slides),'slides':slides},ensure_ascii=False,indent=2)+'\n')
notes=['# Выступление Alashi','', 'План: 5 минут вместе с двухминутным видео. Основные слайды: 1-7. Слайды 8-9 для вопросов.','']
for i,s in enumerate(slides,1):notes += [f'## {i}. {s["title"]}',f'Время: {s["seconds"]} секунд.' if s['seconds'] else 'Запасной слайд.',s['notes'],'',f'Источники: {s["sources"]}','']
(ROOT/'docs/DEMO_TALK_20260906.md').write_text('\n'.join(notes))
md=['# Alashi: презентация демо-дня','', 'Семь основных слайдов и два запасных. Суммарный целевой хронометраж: 5 минут, включая видео 2 минуты.','']
for i,s in enumerate(slides,1):
 md += [f'## {i}. {s["title"]}','']+[n['value'].replace('\n',' ') for n in s['nodes'] if n['kind']=='text']+['']
(ROOT/'docs/DECK.md').write_text('\n'.join(md))

# PPTX: actual editable text and shapes, plus embedded video on slide 4.
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
def measure(v):return Inches(v/144)
def rgb(c):return RGBColor.from_string(c.lstrip('#'))
for data in slides:
 page=prs.slides.add_slide(prs.slide_layouts[6]);page.background.fill.solid();page.background.fill.fore_color.rgb=rgb(BG)
 for n in data['nodes']:
  x,y,w,h=[measure(n[k]) for k in ['x','y','w','h']]
  if n['kind']=='rect':
   sh=page.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h);sh.fill.solid();sh.fill.fore_color.rgb=rgb(n['color']);sh.line.fill.background()
  elif n['kind']=='image':page.shapes.add_picture(n['path'],x,y,w,h)
  elif n['kind']=='video':page.shapes.add_movie(str(VIDEO),x,y,w,h,poster_frame_image=n['path'],mime_type='video/mp4')
  else:
   sh=page.shapes.add_textbox(x,y,w,h);tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.TOP
   for j,line in enumerate(n['value'].split('\n')):
    para=tf.paragraphs[0] if j==0 else tf.add_paragraph();para.text=line;para.space_before=Pt(0);para.space_after=Pt(0);para.line_spacing=1.08
    for run in para.runs:
     run.font.name=n.get('font','Jura');run.font.size=Pt(n['size']/2);run.font.bold=n.get('bold',False);run.font.color.rgb=rgb(n['color'])
     if n.get('url'):run.hyperlink.address=n['url']
 page.notes_slide.notes_text_frame.text=data['notes']+'\n\nИсточники: '+data['sources']
prs.save(OUT/'deck.pptx')

font64=base64.b64encode((ROOT/'assets/fonts/Jura.ttf').read_bytes()).decode()
css='''@font-face{font-family:Jura;src:url(data:font/ttf;base64,FONT)}*{box-sizing:border-box}html,body{margin:0;width:100%;height:100%;overflow:hidden;background:#03080d;color:#edf3ed}.slide{position:absolute;width:1920px;height:1080px;left:0;top:0;background:#09131c;display:none;transform-origin:top left}.slide.active{display:block}.node{position:absolute;white-space:pre-wrap;margin:0;line-height:1.08}.node.text{font-family:Jura,sans-serif}.node.image{object-fit:contain}.node.video{object-fit:contain;background:#000}#controls{position:fixed;bottom:10px;right:16px;display:flex;gap:12px;align-items:center;color:#a9bbc5;font:16px Jura;background:#09131cdd;padding:8px;border-radius:6px}button{color:inherit;background:#142634;border:1px solid #29414f;border-radius:4px;padding:5px 12px;font:inherit;cursor:pointer}.capture #controls{display:none}#notes{position:fixed;left:30px;right:30px;bottom:65px;padding:24px;background:#142634f5;font:23px/1.4 Jura;display:none}a{color:inherit;text-decoration:none}@media print{@page{size:1920px 1080px;margin:0}html,body{overflow:visible;height:auto}.slide{position:relative;display:block!important;left:0!important;top:0!important;transform:none!important;page-break-after:always}#controls,#notes{display:none!important}video{display:none}img.poster{display:block!important}}'''.replace('FONT',font64)
sections=[]
for i,s in enumerate(slides):
 nodes=[]
 for n in s['nodes']:
  style=';'.join(f'{k}:{n[k]}px' for k in ['left','top','width','height'] if k in n)
  style=f'left:{n["x"]}px;top:{n["y"]}px;width:{n["w"]}px;height:{n["h"]}px;'
  if n['kind']=='rect':nodes.append(f'<div class="node" style="{style}background:{n["color"]}"></div>')
  elif n['kind'] in ['image','video']:
   dataurl='data:image/png;base64,'+base64.b64encode(Path(n['path']).read_bytes()).decode()
   if n['kind']=='video':nodes.append(f'<img class="node image poster" style="{style}display:none" src="{dataurl}"><video class="node video" style="{style}" controls preload="metadata" poster="{dataurl}" src="out/alashi-demo.mp4"></video>')
   else:nodes.append(f'<img class="node image" style="{style}" src="{dataurl}">')
  else:
   value=html.escape(n['value']);value=f'<a href="{n["url"]}" target="_blank" rel="noopener">{value}</a>' if n.get('url') else value
   nodes.append(f'<div class="node text" style="{style}font-size:{n["size"]}px;color:{n["color"]};font-weight:{700 if n.get("bold") else 500};font-family:{n.get("font","Jura")},sans-serif">{value}</div>')
 sections.append(f'<section class="slide" aria-label="{html.escape(s["title"])}" data-slide="{i+1}">'+''.join(nodes)+'</section>')
script='''const slides=[...document.querySelectorAll('.slide')],notes=NOTES;let index=Math.max(0,Math.min(slides.length-1,Number(new URLSearchParams(location.search).get('slide')||1)-1));if(location.search.includes('capture'))document.body.classList.add('capture');function fit(){const scale=Math.min(innerWidth/1920,innerHeight/1080);slides.forEach(s=>{s.style.transform=`scale(${scale})`;s.style.left=(innerWidth-1920*scale)/2+'px';s.style.top=(innerHeight-1080*scale)/2+'px'})}function show(i){document.querySelectorAll('video').forEach(v=>v.pause());index=Math.max(0,Math.min(slides.length-1,i));slides.forEach((s,j)=>s.classList.toggle('active',j===index));document.querySelector('#count').textContent=(index+1)+' / '+slides.length;document.querySelector('#notes').textContent=notes[index];fit()}document.querySelector('#next').onclick=()=>show(index+1);document.querySelector('#prev').onclick=()=>show(index-1);document.querySelector('#full').onclick=()=>document.fullscreenElement?document.exitFullscreen():document.documentElement.requestFullscreen();document.querySelector('#show-notes').onclick=()=>{const n=document.querySelector('#notes');n.style.display=n.style.display==='block'?'none':'block'};addEventListener('resize',fit);addEventListener('keydown',e=>{if(e.target.tagName==='VIDEO')return;if(['ArrowRight','PageDown'].includes(e.key)){e.preventDefault();show(index+1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(index-1)}});window.Deck={show,slides};show(index);'''.replace('NOTES',json.dumps([s['notes'] for s in slides],ensure_ascii=False).replace('<','\\u003c'))
page='<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Alashi · Демонстрация</title><style>'+css+'</style><body>'+''.join(sections)+'<div id="notes"></div><nav id="controls"><button id="prev">←</button><span id="count"></span><button id="next">→</button><button id="show-notes">Заметки</button><button id="full">⛶</button></nav><script>'+script+'</script></body></html>'
(ROOT/'docs/deck.html').write_text(page)
shutil.copy2(VIDEO,OUT/'alashi-demo.mp4')
(PACKAGE/'Alashi.html').write_text(page.replace('out/alashi-demo.mp4','alashi-demo.mp4'))
shutil.copy2(VIDEO,PACKAGE/'alashi-demo.mp4');shutil.copy2(OUT/'deck.pptx',PACKAGE/'Alashi.pptx');shutil.copy2(ROOT/'docs/DEMO_TALK_20260906.md',PACKAGE/'Речь.md');shutil.copy2(ROOT/'assets/fonts/Jura.ttf',PACKAGE/'Jura.ttf')
print(json.dumps({'slides':len(slides),'main':7,'seconds':sum(s['seconds'] for s in slides),'pptx_bytes':(OUT/'deck.pptx').stat().st_size,'package':str(PACKAGE)},ensure_ascii=False))
